"""Create an editable scene-GCN diagram from the existing mean diagram.

Run: python docs/design/generate_recsys_diagram_gnn.py
Preserves Baseline, Text, the recommender band and the original canvas.
Requires python-pptx, PyMuPDF and LibreOffice; no model is executed.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from generate_main_diagram import Diagram, color, render
from generate_recsys_diagram import FONT, INK, MUTED, TEAL


def build(presentation):
    d = Diagram(presentation)
    slide = d.slide
    # Replace Graph computation only; preserve its panel and heading.
    keep = {"Rounded Rectangle 5", "TextBox 6", "TextBox 8"}
    for shape in list(slide.shapes):
        x, y = shape.left / 914400, shape.top / 914400
        if x >= 9.35 and 2.0 <= y < 10.37 and shape.name not in keep:
            shape._element.getparent().remove(shape._element)

    d.label("GNN model subtitle", 0.65, 0.93, 20.1, 0.30,
            "DCGN-inspired scene GCN  |  Similarity excludes Title  |  Frozen BGE Large: 1,024",
            16, MUTED, align=PP_ALIGN.LEFT)

    def node(name, x, y, w, h, lines, size=17, fill="FFFFFF"):
        d.box(name, x, y, w, h, fill, TEAL)
        shape = d.text(name + " text", x + 0.07, y + 0.03, w - 0.14, h - 0.06,
                       [(value, size, INK, False) for value in lines])
        for paragraph in shape.text_frame.paragraphs:
            paragraph.space_after = Pt(0)
        return shape

    for role, x, count in (("Title", 9.7, "1,024"), ("Entities", 12.5, "E × 1,024"),
                           ("Actions", 15.3, "A × 1,024"), ("Context", 18.1, "1,024")):
        node("Graph " + role, x, 2.2, 2.5, 0.55, [role], size=18)
        node("Graph " + role + " BGE", x, 2.94, 2.5, 0.60,
             ["BGE [frozen]", count], size=16, fill="EFF3F7")
        node("Graph " + role + " projection", x, 3.75, 2.5, 0.65,
             ["Linear Projection", "1,024 → 128"], size=17)
        d.wire(role + " to BGE", [(x + 1.25, 2.75), (x + 1.25, 2.94)], MUTED)
        d.wire(role + " to projection", [(x + 1.25, 3.54), (x + 1.25, 3.75)], TEAL)

    d.box("Graph Action role fusion", 12.5, 4.66, 5.3, 0.80, "FFFFFF", TEAL)
    fusion = d.text("Graph Action role fusion text", 12.58, 4.70, 5.14, 0.72,
                    [("Entity → Action role fusion", 16, INK, False),
                     ("actor · target · tool · location", 13, MUTED, False),
                     ("Action + ReLU(Linear 640 → 128)", 13, MUTED, False)])
    for paragraph in fusion.text_frame.paragraphs:
        paragraph.space_after = Pt(0)
    d.wire("Entity projection to role fusion", [(13.75, 4.4), (13.75, 4.66)], TEAL)
    d.wire("Action projection to role fusion", [(16.55, 4.4), (16.55, 4.66)], TEAL)
    d.wire("Entity projection to mean", [(13.75, 4.53), (12.28, 4.53),
                                         (12.28, 5.90), (12.5, 5.90)], TEAL)
    d.junction("Entity forward fork", 13.75, 4.53, TEAL)
    node("Graph Entity mean", 12.5, 5.67, 2.5, 0.46, ["Entity mean: 128"], size=15)
    node("Graph Action mean", 15.3, 5.67, 2.5, 0.46, ["Action mean: 128"], size=15)
    d.wire("Enriched Actions to mean", [(16.55, 5.46), (16.55, 5.67)], TEAL)

    node("Graph Content concat", 12.5, 6.38, 8.1, 0.50,
         ["Content X: [Entity mean 128; Action mean 128; Context 128] = 384"], size=15)
    for role, x, y in (("Entities", 13.75, 6.13), ("Actions", 16.55, 6.13),
                       ("Context", 19.35, 4.4)):
        d.wire(role + " to Content concat", [(x, y), (x, 6.38)], TEAL)

    node("Graph Scene concat", 9.7, 7.10, 5.2, 0.50,
         ["Scene concat: [Title 128; Content 384] = 512"], size=14)
    d.wire("Title to Scene concat", [(10.95, 4.4), (10.95, 7.10)], TEAL)
    d.wire("Content to Scene concat", [(16.55, 6.88), (16.55, 7.00),
                                        (14.20, 7.00), (14.20, 7.10)], TEAL)
    d.wire("Content to Similarity", [(16.55, 7.00), (17.95, 7.00), (17.95, 7.10)], TEAL)
    d.junction("Content forward fork", 16.55, 7.00, TEAL)

    node("Graph Scene Residual MLP", 9.7, 7.81, 5.2, 0.62,
         ["Scene Residual MLP → H: S × 512", "512 → 512 → 512 · ReLU · no LN"],
         size=15, fill="DFF3F1")
    d.wire("Concat to Scene MLP", [(12.3, 7.60), (12.3, 7.81)], TEAL)
    d.box("Graph Scene similarity", 15.3, 7.10, 5.3, 1.33, "EFF8F7", TEAL)
    similarity = d.text("Graph Scene similarity text", 15.38, 7.15, 5.14, 1.23,
                        [("Scene similarity → Â: S × S", 17, INK, False),
                         ("384-d X · all pairs within each video", 14, MUTED, False),
                         ("A = max(cos(X), 0); diagonal = 1", 14, INK, False),
                         ("", 16, INK, False)])
    # Native superscript runs keep the powers editable and avoid missing glyphs
    # when the PDF exporter substitutes fonts for Unicode superscript fractions.
    formula = similarity.text_frame.paragraphs[-1]
    formula.clear()
    for value, superscript in (("Â = D", False), ("-1/2", True),
                               (" A D", False), ("-1/2", True)):
        run = formula.add_run()
        run.text = value
        run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(12 if superscript else 16), color(INK)
        if superscript:
            run._r.get_or_add_rPr().set("baseline", "30000")

    node("Graph Scene GCN", 9.7, 8.62, 10.9, 0.65,
         ["Scene GCN · 1 layer · 512 → 512", "H′ = H + ReLU(Â H W + b)"],
         size=17, fill="DFF3F1")
    d.wire("Scene MLP messages to GCN", [(12.3, 8.43), (12.3, 8.62)], TEAL)
    d.wire("Similarity adjacency to GCN", [(17.95, 8.43), (17.95, 8.62)], TEAL)
    node("Graph Mean Pooling", 9.7, 9.47, 10.9, 0.50,
         ["Mean Pooling over valid Scenes: S × 512 → Item 512"], size=17, fill="DFF3F1")
    d.wire("Scene GCN to Mean", [(15.15, 9.27), (15.15, 9.47)], TEAL)
    d.wire("Graph item output", [(15.15, 9.97), (15.15, 10.37)], TEAL)

    d.text("GNN diagram footnotes", 0.65, 12.58, 20.1, 0.52,
           [("E / A = entities / actions per scene; S = valid scenes  |  No scene-order encoding", 13, MUTED, False),
            ("Title-only: existing fallback (GCN skipped)  |  Meta: Baseline  |  Fully missing / padding: 0", 13, MUTED, False)],
           align=PP_ALIGN.LEFT)
    slide.notes_slide.notes_text_frame.text = (
        "Graph experiment: sasrec-role-graph/v8; aggregation=gnn. "
        "A DCGN-inspired recommendation variant, not a reproduction of hierarchical DCGN. "
        "X concatenates entity means, role-enriched action means and projected Context, "
        "before the Scene MLP and without Title. L2 normalization uses eps=1e-8; "
        "cosine weights are clamped at zero, with unit diagonal. "
        "D is the diagonal matrix of row sums of A. Similarity is recomputed each "
        "forward with gradients enabled. GCN messages use the existing 512-dimensional "
        "Scene Residual MLP output. Single scenes receive the self-loop transformation. "
        "No temporal convolution, hierarchy, gate, attention, GCN dropout or LayerNorm. "
        "BGE stays fixed; projections, role fusion, Scene MLP, GCN and SASRec train jointly. "
        "The Baseline, Text path and SASRec retain the existing v7 model structure."
    )
    return presentation


def main():
    directory = Path(__file__).resolve().parent
    output = directory / "recsys_diagram_gnn.pptx"
    build(Presentation(directory / "recsys_diagram.pptx")).save(output)
    render(output, expected_size=(3440, 2112))


if __name__ == "__main__":
    main()
