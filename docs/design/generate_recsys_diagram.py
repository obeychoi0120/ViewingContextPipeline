"""Update the editable recommendation diagram to the implemented v7 model.

Preserves the existing Baseline/Text panels, palette and recommender band.
Run: python docs/design/generate_recsys_diagram.py
Requires python-pptx, PyMuPDF and LibreOffice; no training is executed.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from generate_main_diagram import Diagram, color, render

INK = "19324F"
MUTED = "62768D"
BLUE = "2465C4"
TEAL = "008B91"
GRAY = "AAB1B9"
FONT = "Liberation Sans"


def text(shape, lines, size=18):
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, value in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = value
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color(INK)


def build(presentation):
    d = Diagram(presentation)
    slide = d.slide
    shapes = {s.name: s for s in slide.shapes}
    for name in ("TextBox 2", "TextBox 100"):
        if name in shapes:
            element = shapes[name]._element
            element.getparent().remove(element)

    # Keep the original Baseline/Text shapes and colours; replace only their
    # final transformation and lower connectors. These names remain stable.
    concat = shapes["Rounded Rectangle 116"]
    concat.top, concat.height = Inches(8.10), Inches(0.68)
    text(concat, ["Concat: [Title 128; Summary 384]", "512"], 16)
    for name in ("Rounded Rectangle 160", "Rounded Rectangle 170"):
        mlp = shapes[name]
        mlp.top, mlp.height = Inches(8.95), Inches(0.86)
        text(mlp, ["Item Residual MLP", "512 → 512 → 512", "ReLU · no LayerNorm"], 15)

    # Remove the old Graph computation (or the previous generated version),
    # keeping its panel and heading. Lower Baseline/Text wires are redrawn too.
    keep = {"Rounded Rectangle 5", "TextBox 6", "TextBox 8"}
    for shape in list(slide.shapes):
        x, y = shape.left / 914400, shape.top / 914400
        graph = x >= 9.35 and 2.0 <= y < 10.37 and shape.name not in keep
        graph = graph or shape.name in {"Node-type addition caption", "Entity node-type addition",
                                      "Action node-type addition", "Added node features merge"}
        lower_wire = shape.shape_type == MSO_SHAPE_TYPE.LINE and x < 9.35 and 5.84 <= y < 10.37
        if graph or lower_wire:
            element = shape._element
            element.getparent().remove(element)
    d.wire("Baseline projection to MLP", [(2.15, 5.85), (2.15, 8.95)], GRAY)
    d.wire("Baseline item output", [(2.15, 9.81), (2.15, 10.37)], GRAY)
    for x in (5.445, 7.595):
        d.wire("Text component to concat", [(x, 5.85), (x, 7.65), (6.545, 7.65)], BLUE, arrow=False)
    d.wire("Text concat input", [(6.545, 7.65), (6.545, 8.10)], BLUE)
    d.wire("Text concat to MLP", [(6.545, 8.78), (6.545, 8.95)], BLUE)
    d.wire("Text item output", [(6.545, 9.81), (6.545, 10.37)], BLUE)

    def node(name, x, y, w, h, lines, size=18, fill="FFFFFF"):
        d.box(name, x, y, w, h, fill, TEAL)
        return d.text(name + " text", x + 0.07, y + 0.03, w - 0.14, h - 0.06,
                      [(value, size, INK, False) for value in lines])

    for role, x, count in (("Title", 9.7, "1,024"), ("Entities", 12.5, "E × 1,024"),
                           ("Actions", 15.3, "A × 1,024"), ("Context", 18.1, "1,024")):
        node("Graph " + role, x, 2.2, 2.5, 0.55, [role])
        node("Graph " + role + " BGE", x, 2.98, 2.5, 0.64, ["BGE [frozen]", count], fill="EFF3F7")
        node("Graph " + role + " projection", x, 3.87, 2.5, 0.73,
             ["Linear Projection", "1,024 → 128"])
        d.wire(role + " to BGE", [(x + 1.25, 2.75), (x + 1.25, 2.98)], MUTED)
        d.wire(role + " to projection", [(x + 1.25, 3.62), (x + 1.25, 3.87)], TEAL)

    d.box("Graph Action role fusion", 12.5, 4.96, 5.3, 1.15, "FFFFFF", TEAL)
    d.text("Graph Action role fusion text", 12.58, 5.02, 5.14, 1.03,
           [("Entity → Action role fusion", 18, INK, False),
            ("actor · target · tool · location", 15, MUTED, False),
            ("640 → 128", 15, MUTED, False)])
    d.wire("Entity projection to role fusion", [(13.75, 4.6), (13.75, 4.96)], TEAL)
    d.wire("Action projection to role fusion", [(16.55, 4.6), (16.55, 4.96)], TEAL)
    # Entity vectors stay unchanged; only Actions receive Entity information.
    d.wire("Entity projection to mean", [(13.75, 4.78), (12.28, 4.78),
                                         (12.28, 6.595), (12.5, 6.595)], TEAL)
    d.junction("Entity forward fork", 13.75, 4.78, TEAL)
    node("Graph Entity mean", 12.5, 6.32, 2.5, 0.55, ["Entity mean: 128"], size=16)
    node("Graph Action mean", 15.3, 6.32, 2.5, 0.55, ["Action mean: 128"], size=16)
    d.wire("Enriched Actions to mean", [(16.55, 6.11), (16.55, 6.32)], TEAL)
    node("Graph Scene concat", 9.7, 7.2, 10.9, 0.66,
         ["Scene concat: [Title; Entity mean; Action mean; Context] = 512"], size=18)
    for role, x, y in (("Title", 10.95, 4.6), ("Entities", 13.75, 6.87),
                       ("Actions", 16.55, 6.87), ("Context", 19.35, 4.6)):
        d.wire(role + " to Scene concat", [(x, y), (x, 7.2)], TEAL)
    node("Graph Scene Residual MLP", 9.7, 8.12, 10.9, 0.72,
         ["Scene Residual MLP", "512 → 512 → 512 · ReLU · no LayerNorm"],
         size=17, fill="DFF3F1")
    node("Graph Mean Pooling", 9.7, 9.1, 10.9, 0.62,
         ["Mean Pooling over valid Scenes: S × 512 → Item 512"], size=18, fill="DFF3F1")
    for name, start, end in (("Concat to Scene MLP", 7.86, 8.12),
                             ("Scene MLP to Mean", 8.84, 9.1),
                             ("Graph item output", 9.72, 10.37)):
        d.wire(name, [(15.15, start), (15.15, end)], TEAL)
    text(shapes["TextBox 90"], ["Last valid position", "User Residual MLP · GELU + LN"], 16)
    return presentation


def main():
    output = Path(__file__).resolve().parent / "recsys_diagram.pptx"
    build(Presentation(output)).save(output)
    # Keep the original 21.5 × 13.2 inch canvas and use the main deck's 160 dpi.
    render(output, expected_size=(3440, 2112))


if __name__ == "__main__":
    main()
