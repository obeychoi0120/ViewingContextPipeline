"""Update both editable recommendation diagrams and render PNGs via LibreOffice.

Run from any directory with python-pptx, PyMuPDF and LibreOffice installed.
The large and small diagrams use one node/context shared projection box.
"""

from copy import deepcopy
from pathlib import Path
import os
import shutil
import subprocess
import tempfile

import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt


def remove(shape):
    element = shape._element
    element.getparent().remove(element)


def clone(slide, shape):
    element = deepcopy(shape._element)
    properties = element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    properties.set("id", str(slide.shapes._next_shape_id))
    properties.set("name", "Shared projection branch")
    slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return slide.shapes[-1]


def single_line(shape, text, size):
    paragraphs = list(shape.text_frame.paragraphs)
    paragraphs[0].runs[0].text = text
    paragraphs[0].runs[0].font.size = Pt(size)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)


def shared_projection(presentation):
    slide = presentation.slides[0]
    shapes = list(slide.shapes)
    if any(s.has_text_frame and s.text.startswith("Shared Linear Projection:") for s in shapes):
        return
    node = next(s for s in shapes if s.has_text_frame and s.text.startswith("Shared projection:"))
    context = next(s for s in shapes if s.has_text_frame and s.text.startswith("Projection\n")
                   and s.left == Inches(18.1))
    role = next(s for s in shapes if s.has_text_frame and s.text.startswith("Role Graph Encoder"))
    incoming = next(s for s in shapes if not s.has_text_frame and s.left == Inches(15.15)
                    and s.top == Inches(4.6))
    node.width = Inches(8.1)
    node.text_frame.paragraphs[0].runs[0].text = "Shared Linear Projection: 1,024 → 128"
    node.text_frame.paragraphs[1].runs[0].text = "Same weights; applied separately to Entities / Actions / Context"
    node.text_frame.paragraphs[1].runs[0].font.size = Pt(15)
    remove(context)
    embedding = clone(slide, role)
    embedding.top, embedding.height = Inches(4.75), Inches(0.42)
    single_line(embedding, "Entity / action node-type embedding", 16)
    role.top, role.height = Inches(5.34), Inches(0.42)
    single_line(role, "Role Graph Encoder · 1 layer · hidden 128", 16)
    # Two node outputs receive type embeddings; the Context output follows its
    # existing vertical wire directly to the scene concat.
    for x, top, height in ((13.75, 4.6, 0.15), (16.55, 4.6, 0.15), (15.15, 5.17, 0.17)):
        line = clone(slide, incoming)
        line.left, line.top, line.height = Inches(x), Inches(top), Inches(height)
    remove(incoming)
    for line in shapes:
        if not line.has_text_frame and line.top == Inches(5.61) and line.left in (Inches(13.75), Inches(16.55)):
            line.top, line.height = Inches(5.76), Inches(0.14)


def small_model(presentation):
    slide = presentation.slides[0]
    shapes = list(slide.shapes)
    for shape in shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("1,024", "384").replace("BGE large", "BGE small")
    shapes[1].text_frame.paragraphs[0].runs[0].text = (
        "BGE small-en-v1.5  /  Baseline: Title 512  /  Text & Graph: Title 128 + Video 384"
        "  /  Final LayerNorm: 512  /  Frozen BGE: 384"
    )
    # Remove the Text summary projection and its incoming connector. The existing
    # blue wire now runs directly from the summary's BGE output to the concat.
    projection = next(s for s in shapes if s.has_text_frame and s.left == Inches(6.62)
                      and s.top == Inches(5.05))
    incoming = next(s for s in shapes if not s.has_text_frame and s.left == Inches(7.595)
                    and s.top == Inches(4.4))
    outgoing = next(s for s in shapes if not s.has_text_frame and s.left == Inches(7.595)
                    and s.top == Inches(5.85))
    remove(projection)
    remove(incoming)
    outgoing.top, outgoing.height = Inches(4.4), Inches(3.3)
    shapes[26].text_frame.paragraphs[0].runs[0].text = (
        "E / A = entities / actions; S = valid scenes  |  Node/context projection weights shared"
        "  |  Text summary: direct 384  |  No-title: 128 zeros"
    )


def render(output):
    environment = os.environ.copy()
    program = Path(shutil.which("soffice")).resolve().parent
    environment["LD_LIBRARY_PATH"] = str(program) + ":" + environment.get("LD_LIBRARY_PATH", "")
    with tempfile.TemporaryDirectory(prefix="recsys-small-render-") as work:
        work = Path(work)
        subprocess.run([
            "soffice", f"-env:UserInstallation={work.joinpath('lo-profile').as_uri()}",
            "--headless", "--convert-to", "pdf", "--outdir", str(work), str(output),
        ], check=True, env=environment)
        with pymupdf.open(work / output.with_suffix(".pdf").name) as pdf:
            assert len(pdf) == 1
            pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(
                output.with_suffix(".png")
            )


def main():
    directory = Path(__file__).resolve().parent
    regular = directory / "recsys_diagram.pptx"
    presentation = Presentation(regular)
    shared_projection(presentation)
    title_projection = next(
        shape for shape in presentation.slides[0].shapes
        if shape.has_text_frame and shape.left == Inches(9.7) and shape.top == Inches(3.87)
    )
    title_projection.text_frame.paragraphs[0].runs[0].text = "Linear Projection"
    presentation.save(regular)
    render(regular)
    small_model(presentation)
    small = directory / "recsys_diagram_small.pptx"
    presentation.save(small)
    render(small)


if __name__ == "__main__":
    main()
