"""Render G1--G4 as matching 3840 x 2160 architecture PNGs.

Run: python docs/design/generate_graph_arm_diagrams.py
Requires python-pptx, PyMuPDF and LibreOffice, as do the existing generators.
Only PNGs are retained; editable slides and PDFs are temporary render inputs.
Mirrors SceneRoleEncoder in validation/graph_model.py, SceneAttentionGCN in
validation/scene_attention.py and the extensions in validation/graph_context.py.
No model, extraction, embedding or training is executed.
"""

from pathlib import Path
import shutil
import tempfile

from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from generate_main_diagram import Diagram, render


INK, MUTED = "19324F", "62768D"
TEAL, TEAL_FILL = "008B91", "EFF9F8"
PURPLE, PURPLE_FILL = "7139BE", "F3EEFC"
ORANGE, ORANGE_FILL = "B96B16", "FFF5E6"
GREEN, GREEN_FILL = "28835D", "EFF8F1"
GRAY, GRAY_FILL = "9AA7B5", "EFF2F6"

ARMS = {
    "g1": ("Attention-weighted Scene GNN", False, False),
    "g2": ("Scene GNN + Title-seeded Video Node", True, False),
    "g3": ("Scene GNN + Activity Segments", False, True),
    "g4": ("Scene GNN + Activity Segments + Video Node", True, True),
}


def card(d, name, x, y, w, h, title, *lines, ink=TEAL, fill="FFFFFF", size=16,
         detail_size=13):
    d.box(name, x, y, w, h, fill, ink)
    shape = d.text(name + " text", x + 0.07, y + 0.03, w - 0.14, h - 0.06,
                   [(title, size, INK, True),
                    *((line, detail_size, INK, False) for line in lines)])
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
    return shape


def wire(d, name, points, ink=TEAL, *, dashed=False):
    start = len(d.slide.shapes)
    d.wire(name, points, ink)
    for shape in list(d.slide.shapes)[start:]:
        # Remove the inherited Office connector style (including its shadow).
        style = shape._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}style"
        )
        if style is not None:
            shape._element.remove(style)
        if dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def heading(d, name, x, y, w, text, ink=TEAL):
    d.label(name, x, y, w, 0.32, text, 18, ink, True, PP_ALIGN.LEFT)


def step(d, name, x, y, w, h, title, dimensions, note=None, **style):
    """Put tensor or transformation dimensions in brackets on every processing step."""
    lines = [f"[{dimensions}]"]
    if note:
        lines.append(note)
    return card(d, name, x, y, w, h, title, *lines, **style)


def scene_encoder(d, video_node):
    """The same frozen features, role fusion and residual Scene MLP in all Arms."""
    d.box("Scene encoder panel", 0.55, 2.0, 13.85, 5.78, TEAL_FILL, "B9DFDD")
    heading(d, "Scene encoder heading", 0.85, 2.13, 13.1,
            "1  |  ROLE-AWARE SCENE ENCODER")
    columns = (("Title", 0.95, ""), ("Entities", 4.25, "E × "),
               ("Actions", 7.55, "A × "), ("Context", 10.85, ""))
    for name, x, count in columns:
        card(d, name, x, 2.69, 2.95, 0.43, name, size=17)
        step(d, name + " BGE", x, 3.34, 2.95, 0.57, "BGE (frozen)", count + "1024",
             ink=GRAY, fill=GRAY_FILL, size=15)
        step(d, name + " projection", x, 4.13, 2.95, 0.57, "Linear Projection",
             f"{count}1024 → {count}128", size=15, detail_size=12)
        wire(d, name + " input", [(x + 1.475, 3.12), (x + 1.475, 3.34)], MUTED)
        wire(d, name + " encoding", [(x + 1.475, 3.91), (x + 1.475, 4.13)])

    step(d, "Role fusion", 4.25, 4.95, 6.25, 0.72, "Entity → Action Role Fusion",
         "A × 640 → A × 128", "Residual fusion · actor / target / tool / place",
         size=15, detail_size=12)
    wire(d, "Entity roles", [(5.725, 4.70), (5.725, 4.95)])
    wire(d, "Action roles", [(9.025, 4.70), (9.025, 4.95)])
    wire(d, "Entity bypass", [(5.725, 4.82), (4.03, 4.82), (4.03, 6.06), (4.25, 6.06)])
    d.junction("Entity fork", 5.725, 4.82, TEAL)
    step(d, "Entity mean", 4.25, 5.84, 2.95, 0.48, "Entity Mean", "E × 128 → 128",
         size=14, detail_size=12)
    step(d, "Action mean", 7.55, 5.84, 2.95, 0.48, "Action Mean", "A × 128 → 128",
         size=14, detail_size=12)
    wire(d, "Fused actions", [(9.025, 5.67), (9.025, 5.84)])
    step(d, "Content X", 4.25, 6.48, 9.55, 0.48,
         "Content X: Entity / Action / Context Concat", "S × 384", size=14,
         detail_size=12)
    for name, x, start in (("Entity", 5.725, 6.32), ("Action", 9.025, 6.32),
                            ("Context", 12.325, 4.70)):
        wire(d, name + " content", [(x, start), (x, 6.48)])

    title_slot = "Zero Title slot" if video_node else "Title + Content"
    step(d, "Scene concat", 0.95, 7.15, 6.05, 0.48,
         f"Scene Concat: {title_slot}", "S × 512", size=14, detail_size=12)
    wire(d, "Content to scene", [(5.725, 6.96), (5.725, 7.15)])
    if video_node:
        step(d, "Projected title port", 0.95, 5.18, 2.95, 0.72,
             "Projected Title", "128", "Used in Module 3", ink=ORANGE, fill=ORANGE_FILL,
             size=14, detail_size=12)
        wire(d, "Title port", [(2.425, 4.70), (2.425, 5.18)], ORANGE)
    else:
        wire(d, "Title to scene", [(2.425, 4.70), (2.425, 7.15)])
    step(d, "Scene MLP", 7.55, 7.15, 6.25, 0.48,
         "Scene H: Residual Scene MLP", "S × 512 → S × 512", size=14, detail_size=12)
    wire(d, "Scene concat to MLP", [(7.0, 7.39), (7.55, 7.39)])


def scene_exchange(d):
    d.box("Attention panel", 14.75, 2.0, 8.70, 5.78, PURPLE_FILL, "D8C7F1")
    heading(d, "Attention heading", 15.05, 2.13, 8.1,
            "2  |  ATTENTION-WEIGHTED SCENE GNN", PURPLE)
    step(d, "Attention input", 15.15, 2.69, 4.20, 0.75,
         "Content X", "S × 384", "From Module 1 · title excluded",
         ink=PURPLE, size=16, detail_size=12)
    step(d, "Scene values input", 20.05, 2.69, 3.0, 0.75,
         "Scene H", "S × 512", "From Module 1",
         ink=PURPLE, size=16, detail_size=12)
    step(d, "Attention Q K", 15.15, 3.88, 4.20, 0.75, "Query / Key Projection",
         "384 → 128", "Unit RMS normalization", ink=PURPLE, size=16, detail_size=13)
    wire(d, "X to Q K", [(17.25, 3.44), (17.25, 3.88)], PURPLE)
    step(d, "Attention A", 15.15, 5.07, 4.20, 0.72,
         "Scene Attention Weights", "S × S", "Scaled scores + row softmax",
         ink=PURPLE, size=16, detail_size=13)
    wire(d, "Scores to A", [(17.25, 4.63), (17.25, 5.07)], PURPLE)
    step(d, "Attention GNN", 15.15, 6.18, 7.90, 0.72,
         "Residual Message Passing", "S × 512 → S × 512", "Attention-weighted scene updates",
         ink=PURPLE, size=16, detail_size=13)
    wire(d, "A to GNN", [(17.25, 5.79), (17.25, 6.18)], PURPLE)
    wire(d, "H values to GNN", [(21.55, 3.44), (21.55, 6.18)], PURPLE)
    card(d, "Updated scenes", 15.15, 7.16, 7.90, 0.45,
         "Updated Scenes [S × 512]", ink=PURPLE, size=16)
    wire(d, "GNN updated scenes", [(19.10, 6.90), (19.10, 7.16)], PURPLE)


def output(d, note, *, x=19.85, y=10.04, w=3.20):
    step(d, "Video output", x, y, w, 0.80, "Video Vector", "512", note,
         ink=INK, fill="E8EFF6", size=15, detail_size=12)


def scene_input(d, x=0.95, y=8.83):
    step(d, "Updated scenes input", x, y, 4.20, 0.80,
         "Updated Scenes", "S × 512", "From Module 2", ink=PURPLE,
         fill=PURPLE_FILL, size=15, detail_size=12)


def video_start(d, *, x=0.95, y=8.83, w=4.20):
    step(d, "Video initialization", x, y, w, 0.80,
         "Video Node Initialization", "128 → 512", "Title from Module 1 · pad + Scene MLP",
         ink=ORANGE, fill=ORANGE_FILL, size=14, detail_size=12)


def video_steps(d, x, y, w):
    """Show the three sequential operations with short, aligned connections."""
    gap = 0.24
    width = (w - 2 * gap) / 3
    labels = (("1  Scenes → Video", "S × 512, 512 → 512", "Mean gather + residual"),
              ("2  Video → Scenes", "S × 512, 512 → S × 512", "Broadcast + residual"),
              ("3  Scenes → Video", "S × 512, 512 → 512", "Mean gather + residual"))
    for i, (title, dimensions, note) in enumerate(labels):
        left = x + i * (width + gap)
        step(d, f"Video step {i + 1}", left, y, width, 0.80, title, dimensions, note,
             ink=ORANGE, fill=ORANGE_FILL, size=14, detail_size=12)
        if i:
            wire(d, f"Video step {i} to {i + 1}",
                 [(left - gap, y + 0.40), (left, y + 0.40)], ORANGE)


def aggregation_heading(d, title, ink, *, segments=False, video=False):
    heading(d, "Aggregation heading", 0.85, 8.25, 12.0, "3  |  " + title, ink)
    notes = []
    if segments:
        notes.append("Both segment gates init: 0")
    if video:
        notes.append("Gather 1 & 3 share Linear [1024 → 512]")
    if notes:
        d.label("Aggregation notes", 12.50, 8.28, 10.50, 0.22,
                "  |  ".join(notes), 12, ink)


def g1_readout(d):
    aggregation_heading(d, "G1: MEAN READOUT", PURPLE)
    scene_input(d, y=9.40)
    step(d, "Scene mean readout", 8.20, 9.40, 6.10, 0.80,
         "Mean Pooling", "S × 512 → 512", "Equal weight per valid scene",
         ink=PURPLE, fill=PURPLE_FILL, size=18, detail_size=14)
    wire(d, "Scenes to mean", [(5.15, 9.80), (8.20, 9.80)], PURPLE)
    wire(d, "Mean to output", [(14.30, 9.80), (19.40, 9.80)], PURPLE)
    output(d, "Mean over updated scenes", x=19.40, y=9.40, w=3.65)
    d.label("G1 readout note", 1.0, 10.65, 21.50, 0.25,
            "Attention learns scene-to-scene transfer; final pooling uses a simple mean.",
            14, MUTED, align=PP_ALIGN.LEFT)


def g2_readout(d):
    aggregation_heading(d, "G2: TITLE-SEEDED VIDEO NODE", ORANGE, video=True)
    video_start(d)
    scene_input(d, y=10.04)
    video_steps(d, 5.65, 9.45, 13.65)
    wire(d, "Initial video input", [(5.15, 9.23), (5.40, 9.23),
                                      (5.40, 9.67), (5.65, 9.67)], ORANGE)
    wire(d, "Scene input to video node", [(5.15, 10.44), (5.40, 10.44),
                                           (5.40, 10.03), (5.65, 10.03)], PURPLE)
    wire(d, "Video node output", [(19.30, 9.85), (19.85, 9.85)], ORANGE)
    output(d, "Final video node", y=9.45)


def segment_units(d):
    scene_input(d)
    step(d, "Prepare segments", 0.95, 10.04, 4.20, 0.80,
         "prepare-module", "M × 2", "Same activity key in adjacent scenes",
         ink=GREEN, fill=GREEN_FILL, size=15, detail_size=12)
    step(d, "Scene to segments", 5.65, 8.83, 5.0, 0.80,
         "Scene → Segment Encoding", "K × 514 → K × 512",
         "Scene means + two length features",
         ink=GREEN, fill=GREEN_FILL, size=15, detail_size=13)
    wire(d, "Scenes to segment encoding", [(5.15, 9.23), (5.65, 9.23)], PURPLE)
    wire(d, "Fixed incidence input", [(5.15, 10.44), (5.40, 10.44),
                                         (5.40, 9.43), (5.65, 9.43)], GREEN, dashed=True)


def segment_feedback(d, *, w=7.15):
    step(d, "Segments to scenes", 11.15, 8.83, w, 0.80,
         "Segment → Scene Feedback", "S × 1024 → S × 512",
         "Residual update from incident segments",
         ink=GREEN, fill=GREEN_FILL, size=15, detail_size=12)
    wire(d, "Segment feedback", [(10.65, 9.23), (11.15, 9.23)], GREEN)


def segment_readout(d):
    step(d, "Segment readout", 5.65, 10.04, 5.0, 0.80,
         "Segment Readout", "K × 512 → 512", "Gated mean over segments",
         ink=GREEN, fill=GREEN_FILL, size=15, detail_size=13)
    wire(d, "Segments to readout", [(8.15, 9.63), (8.15, 10.04)], GREEN)
    # One unobstructed skip path below the processing row carries the contribution.
    wire(d, "Segment contribution to output", [(8.15, 10.84), (8.15, 11.09),
                                                  (21.45, 11.09), (21.45, 10.84)], GREEN)


def g3_readout(d):
    aggregation_heading(d, "G3: ACTIVITY SEGMENTS + MEAN", GREEN, segments=True)
    segment_units(d)
    segment_feedback(d)
    segment_readout(d)
    step(d, "Updated scene mean", 11.15, 10.04, 7.15, 0.80,
         "Updated Scene Mean", "S × 512 → 512", "Mean over segment-updated scenes",
         ink=PURPLE, fill=PURPLE_FILL, size=16, detail_size=13)
    wire(d, "Updated scenes to mean", [(14.725, 9.63), (14.725, 10.04)], GREEN)
    wire(d, "Scene mean to output", [(18.30, 10.44), (19.85, 10.44)], PURPLE)
    output(d, "Scene + segment readouts")


def g4_readout(d):
    aggregation_heading(d, "G4: SEGMENTS + VIDEO NODE", GREEN, segments=True, video=True)
    segment_units(d)
    segment_feedback(d, w=4.20)
    segment_readout(d)
    video_start(d, x=15.85, w=3.45)
    video_steps(d, 11.15, 10.04, 8.15)
    wire(d, "Segment scenes to video node", [(12.0, 9.63), (12.0, 10.04)], GREEN)
    wire(d, "Initial video to gather", [(17.575, 9.63), (17.575, 9.84),
                                           (13.10, 9.84), (13.10, 10.04)], ORANGE)
    wire(d, "Video node to output", [(19.30, 10.44), (19.85, 10.44)], ORANGE)
    output(d, "Video + segment readouts")


def recommendation(d):
    d.box("Recommendation panel", 0.55, 11.55, 22.90, 1.40, "F5F7FB", "BCCDDD")
    heading(d, "Recommendation heading", 0.85, 11.65, 12.0,
            "4  |  SHARED SASREC RECOMMENDER", INK)
    d.label("Recommender inputs note", 11.80, 11.67, 6.60, 0.24,
            "History / candidate vectors: from Module 3", 12, MUTED)
    card(d, "Candidate input", 19.05, 11.65, 4.00, 0.33,
         "Candidate Video [512]", ink=INK, fill="E8EFF6", size=13)
    specs = ((0.95, 4.2, "History Video Vectors", "L × 512", "Chronological history; L ≤ 10"),
             (5.65, 4.2, "SASRec", "L × 512 → L × 512", "2 blocks · 2 heads"),
             (10.35, 5.5, "User Readout + MLP", "L × 512 → 512", "Last valid position · GELU + LN"),
             (19.05, 4.0, "Dot Product → Ranking", "512, 512 → 1", "One score per candidate"))
    for i, (x, width, title, dimensions, note) in enumerate(specs):
        step(d, f"Recommendation {i}", x, 12.14, width, 0.65, title, dimensions, note,
             ink=INK, fill="FFFFFF", size=14, detail_size=12)
        if i:
            prev = specs[i - 1]
            wire(d, f"Recommendation flow {i}",
                 [(prev[0] + prev[1], 12.465), (x, 12.465)], MUTED)
    wire(d, "Candidate to scoring", [(21.05, 11.98), (21.05, 12.14)], INK)
    d.label("Training note", 0.65, 13.08, 22.70, 0.25,
            "BGE frozen; other blocks train jointly  |  S: scenes · E / A: entities / actions per scene · "
            "K: segments · M: memberships · L: history length  |  Brackets: feature shapes / dimensions",
            12, MUTED, align=PP_ALIGN.LEFT)


def build(arm):
    title, has_video_node, has_segments = ARMS[arm]
    d = Diagram()
    d.label("Diagram title", 0.60, 0.28, 22.80, 0.60,
            f"{arm.upper()}  |  {title}", 32, INK, True, PP_ALIGN.LEFT)
    subtitle = ("Title seeds the video node; scene title slots are zero." if has_video_node else
                "Title is concatenated into every scene representation.")
    d.label("Diagram subtitle", 0.65, 1.01, 22.7, 0.31,
            "Graph Encoding → Video Representation [512] → SASRec  |  " + subtitle,
            17, MUTED, align=PP_ALIGN.LEFT)
    for name, x, width, ink, fill in (("Frozen BGE", 0.65, 2.55, GRAY, GRAY_FILL),
                                    ("Scene encoder", 3.40, 3.10, TEAL, TEAL_FILL),
                                    ("Scene Attention GNN", 6.70, 4.15, PURPLE, PURPLE_FILL),
                                    ("Video node", 11.05, 2.65, ORANGE, ORANGE_FILL),
                                    ("Activity segments", 13.90, 3.50, GREEN, GREEN_FILL)):
        card(d, "Legend " + name, x, 1.49, width, 0.32, name, ink=ink, fill=fill, size=12)
    d.label("Legend fixed rule note", 17.65, 1.49, 5.60, 0.32,
            "Named inputs link modules; dashed: fixed incidence" if has_segments else
            "Named inputs link modules; solid: feature flow", 11, MUTED)
    d.box("Aggregation panel", 0.55, 8.10, 22.90, 3.15, "FBFCFE", "D6E0E9")
    scene_encoder(d, has_video_node)
    scene_exchange(d)
    {"g1": g1_readout, "g2": g2_readout, "g3": g3_readout, "g4": g4_readout}[arm](d)
    recommendation(d)
    for shape in d.slide.shapes:
        if shape.name.endswith(" panel"):
            shape.adjustments[0] = 0.035
    validate_bounds(d)
    return d.presentation


def validate_bounds(d):
    """Fail before rendering if a block or text box escapes the canvas."""
    for shape in d.slide.shapes:
        if (shape.left < 0 or shape.top < 0
                or shape.left + shape.width > d.presentation.slide_width
                or shape.top + shape.height > d.presentation.slide_height):
            raise ValueError(f"shape outside canvas: {shape.name}")


def main():
    directory = Path(__file__).resolve().parent
    with tempfile.TemporaryDirectory(prefix="graph-arm-diagrams-") as work:
        for arm in ARMS:
            intermediate = Path(work) / f"recsys_diagram_{arm}.pptx"
            build(arm).save(intermediate)
            render(intermediate, expected_size=(3840, 2160))
            target = directory / intermediate.with_suffix(".png").name
            shutil.copyfile(intermediate.with_suffix(".png"), target)
            print(f"Created {target.name} (3840 × 2160)")


if __name__ == "__main__":
    main()
