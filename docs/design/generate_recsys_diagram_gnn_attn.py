"""Create an editable attention-GCN (gnn_attn) diagram from the scene-GCN diagram.

Run: python docs/design/generate_recsys_diagram_gnn_attn.py
Builds on generate_recsys_diagram_gnn.py and keeps its layout; only the scene-weight box,
the GCN box, the subtitle and the footnotes change, so the figure shows gnn_attn alone.
Writes gnn_attn/recsys_diagram_gnn_attn.{pptx,png}; the GNN diagram is not modified.
Requires python-pptx, PyMuPDF and LibreOffice; no model is executed.
Mirrors src/validation/scene_attention.py (SceneAttentionGCN) and graph_context.py.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from generate_main_diagram import Diagram, color, render
from generate_recsys_diagram import FONT, INK, MUTED
from generate_recsys_diagram_gnn import build as build_gnn

PURPLE, PURPLE_FILL = "6D28D9", "F1EBFF"
REPLACED = ("Graph Scene similarity", "Graph Scene GCN", "GNN model subtitle",
            "GNN diagram footnotes")


def rich(d, name, x, y, w, h, paragraphs):
    """Centered text whose paragraphs are runs of (text, size, ink, bold[, "sup" | "sub"])."""
    shape = d.text(name, x, y, w, h, [("", 12, INK, False)] * len(paragraphs))
    for paragraph, runs in zip(shape.text_frame.paragraphs, paragraphs, strict=True):
        paragraph.clear()
        paragraph.space_after = Pt(2)
        for value, size, ink, bold, *shift in runs:
            run = paragraph.add_run()
            run.text = value
            run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
            run.font.color.rgb = color(ink)
            if shift:
                # Native super/subscripts stay editable (see the GNN diagram).
                run._r.get_or_add_rPr().set("baseline", {"sup": "30000", "sub": "-25000"}[shift[0]])
    return shape


def build(presentation):
    d = Diagram(build_gnn(presentation))
    for shape in list(d.slide.shapes):
        if shape.name.startswith(REPLACED):
            shape._element.getparent().remove(shape._element)

    d.label("gnn_attn model subtitle", 0.65, 0.93, 20.1, 0.30,
            "Attention scene GCN (gnn_attn)  |  Scene attention excludes Title  |  "
            "Frozen BGE Large: 1,024", 16, MUTED, align=PP_ALIGN.LEFT)

    # Same geometry as the GNN similarity and GCN boxes; the existing wires still connect.
    d.box("Graph Scene attention", 15.3, 7.10, 5.3, 1.33, PURPLE_FILL, PURPLE)
    rich(d, "Graph Scene attention text", 15.38, 7.14, 5.14, 1.25, [
        [("Scene attention → A: S × S", 17, INK, False)],
        [("384-d X · all pairs within each video", 13, MUTED, False)],
        [("Q = W", 14, INK, False), ("q", 14, INK, False, "sub"), (" û(X),  K = W", 14, INK, False),
         ("k", 14, INK, False, "sub"), (" û(X)   384 → 128", 14, INK, False)],
        [("A = softmax(Q K", 16, INK, False), ("T", 16, INK, False, "sup"),
         (" / √128)", 16, INK, False)],
    ])
    d.box("Graph Scene GCN", 9.7, 8.62, 10.9, 0.65, "DFF3F1", PURPLE)
    rich(d, "Graph Scene GCN text", 9.77, 8.65, 10.76, 0.59, [
        [("Attention Scene GCN · 1 layer · 512 → 512", 17, INK, False)],
        [("H′ = H + ReLU(A H W + b)", 17, INK, False)],
    ])

    d.text("gnn_attn diagram footnotes", 0.65, 12.58, 20.1, 0.52, [
        ("E / A = entities / actions per scene; S = valid scenes  |  No scene-order encoding  |  "
         "û(x) = √384 · x / max(‖x‖, 1e-8): attention inputs only", 13, MUTED, False),
        ("Title-only: existing fallback (attention GCN skipped)  |  Single scene: A = [1]  |  "
         "Meta: Baseline  |  Fully missing / padding: 0", 13, MUTED, False),
    ], align=PP_ALIGN.LEFT)
    d.slide.notes_slide.notes_text_frame.text = (
        "Graph experiment: sasrec-role-graph/v12; aggregation=gnn_attn. "
        "X concatenates entity means, role-enriched action means and projected Context, "
        "before the Scene MLP and without Title. Single-head scaled dot-product attention "
        "over all scenes of a video (self included): Q and K are 384 -> 128 without bias, "
        "initialized N(0, 1/(2 fan_in)), and read û(X) = sqrt(384) X / max(||X||, 1e-8). "
        "A is a row softmax recomputed each forward with gradients enabled. "
        "The residual GCN H' = H + ReLU(A H W + b) uses the 512-dimensional Scene Residual MLP "
        "output, then valid scenes are mean pooled. BGE stays fixed; projections, role fusion, "
        "Scene MLP, attention, GCN and SASRec train jointly. "
        "Text uses sasrec-content-v8; SASRec and its User MLP are unchanged."
    )
    return presentation


def main():
    directory = Path(__file__).resolve().parent
    output = directory / "gnn_attn" / "recsys_diagram_gnn_attn.pptx"
    output.parent.mkdir(exist_ok=True)
    build(Presentation(directory / "recsys_diagram.pptx")).save(output)
    render(output, expected_size=(3440, 2112))


if __name__ == "__main__":
    main()
