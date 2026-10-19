"""Build the editable pipeline overview and render its PNG from the PPTX.

Run: python docs/design/generate_main_diagram.py
Requires python-pptx, PyMuPDF and LibreOffice. No recommendation model is run.
The overview includes item representations and the shared recommendation flow.
"""

import os
from pathlib import Path
import shutil
import subprocess
import tempfile

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


FONT = "Liberation Sans"
NAVY = "183450"
MUTED = "536B82"
GRAY = "687581"
BLUE = "2C66B7"
TEAL = "00898D"
WHITE = "FFFFFF"


def color(value):
    return RGBColor.from_string(value)


class Diagram:
    def __init__(self, presentation=None):
        if presentation is not None:
            self.presentation = presentation
            self.slide = presentation.slides[0]
            return
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(24)
        self.presentation.slide_height = Inches(13.5)
        self.slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = color(WHITE)

    def box(self, name, x, y, w, h, fill=WHITE, stroke="CCD7E2"):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.name = name
        style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            shape._element.remove(style)
        shape.adjustments[0] = 0.12
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(1.2)
        return shape

    def text(self, name, x, y, w, h, lines, align=PP_ALIGN.CENTER):
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = Inches(0.025)
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for index, (value, size, ink, bold) in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(4 if index < len(lines) - 1 else 0)
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(ink)
        return shape

    def label(self, name, x, y, w, h, value, size=22, ink=MUTED, bold=False,
              align=PP_ALIGN.CENTER):
        return self.text(name, x, y, w, h, [(value, size, ink, bold)], align=align)

    def node(self, name, x, y, w, h, title, subtitle=None, ink=NAVY,
             fill=WHITE, size=24, subtitle_size=22):
        self.box(name, x, y, w, h, fill, ink)
        lines = [(title, size, ink, True)]
        if subtitle:
            lines.append((subtitle, subtitle_size, MUTED, False))
        return self.text(name + " text", x + 0.10, y + 0.06, w - 0.20, h - 0.12, lines)

    def wire(self, name, points, ink=MUTED, arrow=True):
        for index, (start, end) in enumerate(zip(points, points[1:])):
            line = self.slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]),
            )
            line.name = f"{name} segment {index + 1}"
            line.line.color.rgb = color(ink)
            line.line.width = Pt(1.8)
            properties = line.line._get_or_add_ln()
            if arrow and index == len(points) - 2:
                tail = OxmlElement("a:tailEnd")
                tail.set("type", "triangle")
                tail.set("w", "sm")
                tail.set("len", "sm")
                properties.append(tail)

    def junction(self, name, x, y, ink=MUTED):
        dot = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x - 0.035), Inches(y - 0.035),
            Inches(0.07), Inches(0.07),
        )
        dot.name = name
        dot.fill.solid()
        dot.fill.fore_color.rgb = color(ink)
        dot.line.fill.background()


def build():
    d = Diagram()
    d.box("Shared extraction panel", 0.6, 2.15, 5.0, 9.65, "F6F8FB")
    d.box("Baseline panel", 6.15, 2.15, 11.4, 1.5, "F5F6F8")
    d.box("Text path panel", 6.15, 4.15, 11.4, 3.65, "F0F5FF", "C6D9F5")
    d.box("Graph path panel", 6.15, 8.15, 11.4, 3.65, "EFF9F8", "BADEDC")
    d.label("Title", 0.6, 0.35, 22.8, 0.65, "ViewingContextPipeline", 40, NAVY,
            True, PP_ALIGN.LEFT)
    d.label("Dataset", 0.6, 1.12, 22.8, 0.4, "MicroLens-100K", 24, MUTED,
            align=PP_ALIGN.LEFT)
    for x, w, text in ((0.6, 5.0, "VIDEO PROCESSING"),
                       (6.15, 11.4, "ITEM REPRESENTATIONS")):
        d.label(text, x, 1.63, w, 0.42, text, 28, NAVY, True, PP_ALIGN.LEFT)

    # Scene Graphs feed both paths; Descriptions feed only video-wise Summary.
    d.node("Video", 1.1, 2.85, 4.0, 0.8, "Video")
    d.node("Keyframes", 1.1, 4.0, 4.0, 1.1, "Keyframes",
           "1 Scene: 30s, 6 keyframes", subtitle_size=20)
    d.node("Extraction", 1.1, 5.5, 4.0, 1.0, "Scene Extraction",
           "Qwen3-VL · Qwen3.5 · Gemini", subtitle_size=16)
    d.node("Descriptions", 1.1, 7.0, 4.0, 0.9, "Scene Descriptions", ink=BLUE)
    d.node("Graphs", 1.1, 8.9, 4.0, 0.9, "Scene Graphs", ink=TEAL)
    d.wire("Video to keyframes", [(3.1, 3.65), (3.1, 4.0)])
    d.wire("Keyframes to extraction", [(3.1, 5.1), (3.1, 5.5)])
    d.wire("Extraction to descriptions", [(3.1, 6.5), (3.1, 7.0)], BLUE)
    d.wire("Extraction to graphs", [(3.1, 6.5), (3.1, 6.72),
                                     (0.86, 6.72), (0.86, 9.35), (1.1, 9.35)], TEAL)

    # Each path has its own title branch; no optional-title trunks or dashed wires.
    d.node("Baseline title", 6.4, 2.4, 3.6, 1.0, "Title", ink=GRAY)
    d.node("Baseline embedding", 10.35, 2.4, 3.7, 1.0, "BGE [frozen]", "1,024",
           ink=GRAY)
    d.node("Baseline", 14.75, 2.4, 2.55, 1.0, "Projection + MLP", "Item 512",
           ink=GRAY, size=18)
    d.wire("Baseline title to embedding", [(10.0, 2.9), (10.35, 2.9)], GRAY)
    d.wire("Baseline embedding to output", [(14.05, 2.9), (14.75, 2.9)], GRAY)

    # Independently project frozen 1024-dimensional title and Summary features.
    d.label("Text heading", 6.4, 4.4, 3.0, 0.42, "TEXT PATH", 28, BLUE, True,
            PP_ALIGN.LEFT)
    d.node("Text title", 6.4, 4.95, 3.6, 1.0, "Title", ink=BLUE)
    d.node("Text title embedding", 10.35, 4.95, 3.7, 1.0, "BGE + Projection", "1,024 → 128",
           ink=BLUE, size=22)
    d.node("Video summary", 6.4, 6.5, 3.6, 1.0, "Video-wise Summary", ink=BLUE,
           size=22)
    d.node("Summary embedding", 10.35, 6.5, 3.7, 1.0, "BGE + Projection", "1,024 → 384",
           ink=BLUE, size=22)
    d.node("Text fusion", 14.75, 5.75, 2.55, 1.3, "Concat + MLP", "Item 512", ink=BLUE,
           fill="E2EEFF", size=20)
    d.wire("Descriptions to summary", [(5.1, 7.45), (5.67, 7.45),
                                         (5.67, 6.7), (6.4, 6.7)], BLUE)
    d.wire("Graphs to summary", [(5.1, 9.35), (5.9, 9.35),
                                   (5.9, 7.2), (6.4, 7.2)], TEAL)
    d.wire("Text title to embedding", [(10.0, 5.45), (10.35, 5.45)], BLUE)
    d.wire("Summary to embedding", [(10.0, 7.0), (10.35, 7.0)], BLUE)
    d.wire("Text title embedding to merge", [(14.05, 5.45), (14.45, 5.45),
                                             (14.45, 6.4)], BLUE, arrow=False)
    d.wire("Summary embedding to merge", [(14.05, 7.0), (14.45, 7.0),
                                          (14.45, 6.4)], BLUE, arrow=False)
    d.junction("Text embedding merge", 14.45, 6.4, BLUE)
    d.wire("Text embeddings to fusion", [(14.45, 6.4), (14.75, 6.4)], BLUE)

    # The projected title joins role-enriched scene features BEFORE scene MLP
    # and the selected scene aggregation. No separate Item MLP follows pooling.
    d.label("Graph heading", 6.4, 8.4, 3.0, 0.42, "GRAPH PATH", 28, TEAL, True,
            PP_ALIGN.LEFT)
    d.node("Graph title", 6.4, 8.95, 3.6, 1.0, "Title", ink=TEAL)
    d.node("Graph title embedding", 10.35, 8.95, 3.7, 1.0, "BGE + Projection", "1,024 → 128",
           ink=TEAL, size=22)
    d.node("Graph elements", 6.4, 10.45, 3.6, 1.0, "Graph Elements",
           "Entities · Actions · Context", ink=TEAL, subtitle_size=20)
    d.box("Graph embedding", 10.35, 10.3, 3.7, 1.3, WHITE, TEAL)
    d.text("Graph embedding text", 10.45, 10.36, 3.5, 1.18,
           [("Scene Features", 24, TEAL, True),
            ("BGE + independent projections", 17, MUTED, False),
            ("Role fusion E → A · 384", 18, MUTED, False)])
    d.node("Graph scene MLP", 14.75, 9.45, 2.55, 0.9, "Scene MLP", "Concat 512", ink=TEAL,
           fill="DFF2F0", size=22, subtitle_size=19)
    d.node("Graph aggregation", 14.75, 10.65, 2.55, 0.8, "Aggregation", "Item 512",
           ink=TEAL, fill="DFF2F0", size=20, subtitle_size=18)
    d.wire("Graphs to graph elements", [(5.9, 9.35), (5.9, 10.95), (6.4, 10.95)], TEAL)
    d.junction("Graph source fork", 5.9, 9.35, TEAL)
    d.wire("Graph title to embedding", [(10.0, 9.45), (10.35, 9.45)], TEAL)
    d.wire("Graph elements to embedding", [(10.0, 10.95), (10.35, 10.95)], TEAL)
    d.wire("Graph title embedding to merge", [(14.05, 9.45), (14.45, 9.45),
                                              (14.45, 9.9)], TEAL, arrow=False)
    d.wire("Graph video embedding to merge", [(14.05, 10.95), (14.45, 10.95),
                                              (14.45, 9.9)], TEAL, arrow=False)
    d.junction("Graph scene input merge", 14.45, 9.9, TEAL)
    d.wire("Graph features to scene MLP", [(14.45, 9.9), (14.75, 9.9)], TEAL)
    d.wire("Scene MLP to aggregation", [(16.025, 10.35), (16.025, 10.65)], TEAL)

    # Alternative per-path models, not concatenation of their final outputs.
    d.wire("Baseline item vectors", [(17.3, 2.9), (17.8, 2.9)], GRAY, arrow=False)
    d.wire("Text item vectors", [(17.3, 6.4), (17.8, 6.4)], BLUE, arrow=False)
    d.wire("Graph item vectors", [(17.3, 11.05), (17.8, 11.05)], TEAL, arrow=False)
    d.wire("Alternative item vector bus", [(17.8, 2.9), (17.8, 11.05)], NAVY,
           arrow=False)
    for y in (2.9, 6.4, 7.4, 11.05):
        d.junction("Item bus junction", 17.8, y, NAVY)
    d.box("Recommendation", 18.55, 4.0, 4.4, 7.8, "E7EDF5", NAVY)
    d.label("Recommendation heading", 18.75, 4.3, 4.0, 0.6,
            "Recommendation", 24, NAVY, True)
    d.text("User representation", 18.85, 5.15, 3.8, 1.85,
           [("Viewing history", 22, NAVY, True),
            ("Past item vectors", 19, MUTED, False),
            ("SASRec + User MLP", 19, NAVY, False),
            ("User 512", 20, NAVY, True)])
    d.text("Candidate scoring", 18.85, 7.65, 3.8, 1.85,
           [("Candidate scoring", 22, NAVY, True),
            ("Candidate item vectors", 18, MUTED, False),
            ("User · Candidate", 20, NAVY, True),
            ("Full-catalog ranking", 19, MUTED, False)])
    d.text("Evaluation", 18.85, 10.05, 3.8, 1.25,
           [("Evaluation", 22, NAVY, True),
            ("Validation: HR@10", 19, MUTED, False),
            ("Test: NDCG / HR", 19, MUTED, False)])
    d.label("Aggregation choices", 0.6, 12.18, 22.8, 0.5,
            "Graph aggregation: mean · gnn · attn_pool · gnn_attn · transformer · "
            "transformer_pos · transformer_local", 19, TEAL, align=PP_ALIGN.LEFT)
    d.wire("Item vectors to recommendation", [(17.8, 7.4), (18.55, 7.4)], NAVY)
    d.slide.notes_slide.notes_text_frame.text = (
        "Scene extraction supports qwen3vl, qwen3.5 and gemini. "
        "Graphs feed both paths; Descriptions feed only the Summary-based Text path. "
        "Title is excluded from extraction and summarization. Title branches shown here "
        "are used only by title-enabled arms; other arms mask the title component. "
        "BGE features are frozen; video encoders and SASRec train jointly. "
        "The three item-vector paths are alternative models, not concatenated outputs. "
        "Direct Graph uses the selected aggregation after its Scene MLP. "
        "Current arm names and required inputs are listed in README."
    )
    return d.presentation


def render(output, *, expected_size=(3840, 2160)):
    environment = os.environ.copy()
    program = Path(shutil.which("soffice")).resolve().parent
    environment["LD_LIBRARY_PATH"] = str(program) + ":" + environment.get("LD_LIBRARY_PATH", "")
    with tempfile.TemporaryDirectory(prefix="main-diagram-render-") as work:
        work = Path(work)
        subprocess.run([
            "soffice", f"-env:UserInstallation={work.joinpath('lo-profile').as_uri()}",
            "--headless", "--convert-to", "pdf", "--outdir", str(work), str(output),
        ], check=True, env=environment)
        with pymupdf.open(work / output.with_suffix(".pdf").name) as pdf:
            assert len(pdf) == 1
            page = pdf[0]
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(160 / 72, 160 / 72), alpha=False)
            assert (pixmap.width, pixmap.height) == expected_size
            pixmap.save(output.with_suffix(".png"))


def main():
    output = Path(__file__).resolve().parent / "main_diagram.pptx"
    build().save(output)
    render(output)


if __name__ == "__main__":
    main()
